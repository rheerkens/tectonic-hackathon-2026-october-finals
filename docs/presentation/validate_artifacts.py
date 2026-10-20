"""Document package checks and a preview contact sheet. Does not run app tests."""
from pathlib import Path
import sys, json, re, hashlib, zipfile, io
from xml.etree import ElementTree as E
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'.local/presentation/vendor'))
from pptx import Presentation
from PIL import Image, ImageDraw
from fontTools.ttLib import TTFont
import pymupdf

folder=ROOT/'docs/presentation'
path=folder/'SDtrust_Finals_Draft_v1.pptx'
p=Presentation(path)
visible='\n'.join(s.text for pg in p.slides for s in pg.shapes if s.has_text_frame)
notes=[pg.notes_slide.notes_text_frame.text for pg in p.slides]
native=json.loads((folder/'final_native_validation.json').read_text(encoding='utf-8-sig'))
with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    font_parts=[n for n in z.namelist() if n.endswith('.fntdata')]
    media_parts=[n for n in z.namelist() if n.startswith('ppt/media/')]
    external_media=[]
    for n in z.namelist():
        if n.endswith('.rels'):
            for rel in E.fromstring(z.read(n)):
                if rel.get('TargetMode')=='External' and rel.get('Type','').endswith(('/image','/audio','/video')):
                    external_media.append(rel.attrib)
    embedded=z.read('ppt/presentation.xml').decode()
assert len(p.slides)==8 and len(notes)==8
assert all('EVIDENCE BOUNDARY' in n and 'SOURCES' in n for n in notes)
assert 'No message is sent.' in visible and 'DEMO PLACEHOLDER' in visible
assert re.search('[\u2010-\u2015-]',visible) is None
assert native['overflowCount']==0 and native['opened']
assert len(font_parts)==2 and 'saveSubsetFonts="0"' in embedded
assert not external_media
pdf=pymupdf.open(folder/'SDtrust_Finals_Draft_v1.pdf')
assert len(pdf)==8
fonts=sorted({f[3] for pg in pdf for f in pg.get_fonts()})
# Office renames the embedded families internally to 1 and 2 on PDF export.
# Check the actual font data, not just those PDF resource names.
font_evidence=[]
seen=set()
for pg in pdf:
    for resource in pg.get_fonts():
        if resource[0] in seen: continue
        seen.add(resource[0])
        _,_,_,data=pdf.extract_font(resource[0])
        extracted=TTFont(io.BytesIO(data))
        name=extracted['name'].getDebugName(6)
        assert name in ['Manrope-Regular','Manrope-Bold'], name
        original=TTFont(folder/('assets/'+name+'.ttf'))
        checked=[]
        for char in 'ASDtru012':
            ec=extracted.getBestCmap(); oc=original.getBestCmap()
            if ord(char) not in ec: continue
            eg=ec[ord(char)]; og=oc[ord(char)]
            assert extracted['hmtx'][eg][0] == original['hmtx'][og][0]
            assert extracted['glyf'][eg].getCoordinates(extracted['glyf'])[0] == original['glyf'][og].getCoordinates(original['glyf'])[0]
            checked.append(char)
        assert checked
        font_evidence.append({'pdf_resource':resource[3],'postscript_name':name,'glyph_outlines_and_widths_match':checked})
for pg in pdf:
    assert abs(pg.rect.width / pg.rect.height -16/9)<0.001
preview=folder/'preview'
files=sorted(preview.glob('slide_*.png'))
assert len(files)==8
canvas=Image.new('RGB',(1600,1880),'#e5e1d8'); draw=ImageDraw.Draw(canvas)
for i,f in enumerate(files):
    im=Image.open(f); assert im.size==(1600,900)
    x=16+(i%2)*800; y=38+(i//2)*470
    canvas.paste(im.resize((768,432)),(x,y))
    draw.text((x,y-22),f'Slide {i+1}',fill='#192d46')
canvas.save(preview/'contact_sheet.png')
report={
    'slides':8,'aspect_ratio':'16:9','native_reopen':native['opened'],
    'measured_text_overflows':native['overflowCount'],'notes_with_evidence_boundaries':8,
    'embedded_font_parts':font_parts,'pdf_fonts':fonts,'pdf_font_evidence':font_evidence,'embedded_media':media_parts,
    'external_media':external_media,'visible_dashes':0,'preview_dimensions':[1600,900],
    'script_words':sum(len(m['spoken'].split()) for m in json.loads((folder/'slide_content.json').read_text(encoding='utf-8'))),
    'planned_seconds':170,
    'limits':['No app operation or app build performed.','No app tests run.','Demo is an editable placeholder.','Jury device and live demo playback not verified.','Pitch timing is a rehearsal budget, not a measured spoken performance.'],
    'artifacts':{}
}
for f in [path,folder/'SDtrust_Finals_Draft_v1.pdf',folder/'SDtrust_Finals_Draft_v1_Notes.md',preview/'contact_sheet.png']:
    report['artifacts'][f.name]={'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()}
(folder/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
