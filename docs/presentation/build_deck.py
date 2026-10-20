"""Build the editable first pitch draft. Presentation assets only; no app imports."""
from pathlib import Path
import sys, json, shutil, hashlib, re

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/presentation'
TMP = ROOT / '.local/presentation'
sys.path.insert(0, str(TMP / 'vendor'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from PIL import Image
from fontTools.ttLib import TTFont
from fontTools.varLib.instancer import instantiateVariableFont

ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
sources = {
    'logo.svg': 'apps/web/public/sdtrust.svg',
    'Finn.png': 'docs/design/finn.png',
    'Finn_wave.png': 'docs/design/finn_pet_refined/approved/waving03/frames/03.png',
    'Manrope.ttf': 'apps/web/public/fonts/Manrope.ttf',
    'OFL.txt': 'apps/web/public/fonts/OFL.txt',
}
for dest, source in sources.items():
    shutil.copy2(ROOT / source, ASSETS / dest)

# Static instances improve PowerPoint font support. Keep the upstream font and licence.
for weight, style in [(400, 'Regular'), (700, 'Bold')]:
    font = TTFont(ASSETS / 'Manrope.ttf')
    font = instantiateVariableFont(font, {'wght': weight}, inplace=True)
    names = {1: 'Manrope', 2: style, 3: 'Manrope '+style+' SDtrust',
             4: 'Manrope '+style, 6: 'Manrope-'+style, 16: 'Manrope', 17: style}
    for nid, value in names.items():
        for plat, enc, lang in [(3, 1, 0x409), (1, 0, 0)]:
            font['name'].setName(value, nid, plat, enc, lang)
    font['OS/2'].usWeightClass = weight
    font['OS/2'].fsSelection = (1 << 5) if weight == 700 else (1 << 6)
    font['head'].macStyle = 1 if weight == 700 else 0
    font.save(ASSETS / ('Manrope-'+style+'.ttf'))

NAVY='192D46'; IVORY='F8F6F1'; CORAL='EE6558'; WHITE='FFFEFD'
MUTED='5F6F85'; PALE='E9EEF3'; BORDER='E5E1D8'; AMBER='A76400'
GREEN='0F7A3A'; LIGHT_GREEN='EAF3EC'; AMBER_BG='FFF0D7'
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'SDtrust | Tectonic Hackathon Finals | Draft v1'
prs.core_properties.subject = 'An evidence grounded pitch with an editable demo storyboard'
prs.core_properties.author = 'Sebastian, Roy and Wanne'
prs.core_properties.keywords = 'SDtrust, Finn, SD Worx, Tectonic, draft'
prs.core_properties.comments = 'Synthetic fixture examples. This deck does not certify app runtime or security.'

def rgb(hexcode): return RGBColor.from_string(hexcode)
def rect(s,x,y,w,h,fill,line=None,radius=False):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    if radius:
        q.adjustments[0]=0.10
    q.fill.solid(); q.fill.fore_color.rgb=rgb(fill)
    if line: q.line.color.rgb=rgb(line); q.line.width=Pt(1)
    else: q.line.fill.background()
    return q

def text(s,txt,x,y,w,h,size=24,color=NAVY,bold=False,align=None):
    assert not re.search('[\u2010-\u2015-]',txt), repr(txt)
    q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    q.name='Text: '+txt.replace('\n',' ')[:70]
    tf=q.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=0
    tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Manrope'; p.font.size=Pt(size)
        p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=1.08
        if align is not None: p.alignment=align
        p._p.get_or_add_pPr().set('marL','0')
        p._p.get_or_add_pPr().set('indent','0')
    return q

def picture(s,file,x,y,w,h):
    im=Image.open(file); iw,ih=im.size
    scale=min(w/iw,h/ih)
    dw,dh=iw*scale,ih*scale
    q=s.shapes.add_picture(str(file), Inches(x+(w-dw)/2), Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))
    q.name='Approved asset: '+Path(file).name
    return q

meta=[]
def slide(title,section,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(NAVY if dark else IVORY)
    # Official SVG logo is added by the native PowerPoint render step.
    text(s,'SDtrust',1.17,0.38,2.1,0.44,24,IVORY if dark else NAVY,True)
    text(s,section.upper(),9.05,0.49,3.6,0.28,11,'D6DEE8' if dark else MUTED)
    text(s,'Tectonic Hackathon · SD Worx · Draft v1',0.68,7.02,10.3,0.25,11,'D6DEE8' if dark else MUTED)
    text(s,f'{len(prs.slides):02}',12.2,6.99,0.5,0.28,12,'D6DEE8' if dark else MUTED,True)
    if title: text(s,title,0.7,1.25,11.9,1.18,36,IVORY if dark else NAVY,True)
    return s

def notes(s,title,seconds,spoken,cues,status,refs):
    content=f'{title}\nTiming: {seconds} seconds.\n\nSPEAK\n{spoken}\n\nPRESENTER CUES\n{cues}\n\nEVIDENCE BOUNDARY\n{status}\n\nSOURCES\n'+'\n'.join(refs)
    s.notes_slide.notes_text_frame.text=content
    meta.append(dict(slide=len(prs.slides),title=title,seconds=seconds,spoken=spoken,cues=cues,status=status,sources=refs))

s=slide(None,'The idea',True)
rect(s,0.7,1.42,0.78,0.06,CORAL)
text(s,'The answer.\nAnd the right\nnext step.',0.7,1.8,8.1,2.7,47,IVORY,True)
text(s,'Meet Finn, your SDtrust knowledge companion.',0.74,5.0,8.1,0.8,25,'D6DEE8')
text(s,'Sebastian · Roy · Wanne',0.74,6.15,8,0.42,18,'D6DEE8')
picture(s,ASSETS/'Finn.png',9.1,2.03,3.5,3.92)
notes(s,'The answer and the right next step',15,
    'An HR question should not become a hunt through Teams, emails and client records. Meet Finn, the SDtrust desktop companion for an answer you can check, or a clear next step.',
    'Open calmly. Introduce the team without adding a separate team slide.',
    'Product proposition. The approved Finn artwork and desktop surfaces exist. No live app was operated to create this presentation.',
    ['README.md: product intent and product surfaces','Finals Guide, page 10: fragmented knowledge and calm technology','docs/design/finn.png'])

s=slide('One client. Four different deadlines.','The problem')
text(s,'When can Atlas submit payroll changes?',0.73,2.28,11.8,0.55,25,MUTED)
for x,big,caption,ref in [(0.73,'20','General rule','S1 · Belgium'),(3.9,'22','Client agreement','S4 · Atlas'),(7.07,'23','Meeting proposal','S10 · Teams'),(10.24,'25','Chat recollection','S3 · Teams')]:
    rect(s,x,3.15,2.36,2.25,WHITE,BORDER,True)
    text(s,big,x+0.2,3.42,1.95,0.82,46,NAVY,True)
    text(s,caption,x+0.2,4.42,1.97,0.45,16,NAVY,True)
    text(s,ref,x+0.2,5.0,1.97,0.3,13,MUTED)
text(s,'Finding a message is easy. Knowing what applies is the hard part.',0.74,5.92,11.75,0.66,25,NAVY,True)
notes(s,'Fragmented knowledge',20,
    'In our synthetic Atlas case, a procedure says the twentieth. A client agreement says the twenty second. Teams mentions the twenty third and twenty fifth. Which source actually applies to this client, country and month?',
    'These are October 2026 fixture dates, not competing real SD Worx policies. Do not call every source an approved rule.',
    'Scripted fixture content. No measured claim about how much time employees currently spend searching.',
    ['apps/web/src/state/finn-scenarios.ts: loonmutatieSources, S1, S4, S10, S3','Finals Guide, page 10: knowledge loses context'])

s=slide('Ask once. See the answer and why.','Meet Finn')
picture(s,ASSETS/'Finn_wave.png',0.86,2.7,2.65,3.04)
for yy,num,head,body in [(2.66,'01','Ask in your own words','One small companion on your desktop.'),(3.83,'02','Check the evidence','The answer, its sources and the reasons.'),(5.0,'03','Know what comes next','If evidence is missing, surface the owner.')]:
    text(s,num,4.0,yy+0.03,0.65,0.5,24,CORAL,True)
    text(s,head,4.85,yy,7.1,0.44,26,NAVY,True)
    text(s,body,4.85,yy+0.51,7.4,0.5,21,MUTED)
notes(s,'Meet Finn',15,
    'Finn keeps the interaction small. Ask, see the answer, and open the evidence when needed. If evidence is missing, the experience changes from answering to finding the person who can help.',
    'The question matcher supports the prepared demo questions. Do not suggest arbitrary question coverage. Owner lookup and contact are fixture data.',
    'Prototype source inspected: desktop UI and fixture playback. Full natural language understanding and live search are not demonstrated by this flow.',
    ['apps/web/src/state/finn.ts: match, playFixture, startQuestion','apps/web/src/components/FinnResultContent.tsx: evidence and next step','docs/design/finn_pet_refined/session_status.json: approved waving body_lock_03'])

s=slide('Let’s ask Finn.','Product demo')
rect(s,0.73,2.44,8.15,3.92,PALE,None,True)
rect(s,1.02,2.73,1.55,0.39,NAVY,None,True)
text(s,'DEMO PLACEHOLDER',1.15,2.83,1.31,0.23,8,WHITE,True)
text(s,'Final app capture goes here',1.22,3.72,7.1,0.65,29,NAVY,True)
text(s,'45 seconds · ask, inspect, follow up',1.22,4.52,7.05,0.52,21,MUTED)
text(s,'Editable storyboard. This is not an app screenshot.',1.22,5.68,7.1,0.3,13,MUTED)
for yy,label,body in [(2.54,'ASK','Atlas deadline?'),(3.89,'INSPECT','22 October\nSource S4'),(5.24,'FOLLOW UP','Next month?')]:
    text(s,label,9.35,yy,3.0,0.3,12,CORAL,True)
    text(s,body,9.35,yy+0.39,3.0,0.88,21,NAVY,True)
text(s,'Scripted fixture demo · synthetic records',0.75,6.58,11.7,0.26,13,MUTED)
notes(s,'Product demo',45,
    'This first prototype uses scripted questions and synthetic records. We ask when Atlas can submit payroll changes. Finn shows the twenty second of October and the approved client agreement. We open the evidence to see why the other dates do not apply. Then we ask: and for next month?',
    '0 to 10 seconds: ask “Until when can Atlas submit payroll changes?”\n10 to 25 seconds: show 22 October and open View details or Why this source.\n25 to 35 seconds: point to S4 and one rejected source.\n35 to 45 seconds: ask “And for next month?” and show the evidence gap.\nReplace this placeholder with a final verified recording or current screenshot before submission. If the jury device has no working demo, use slides 5 and 6 as the editable fallback. The guide warns that the jury uses its own device. No video or live link is embedded in v1.',
    'Scripted fixture flow, not live connector search. The inspected frontend matches questions and replays in memory. The README describes a broader future scripted agent with real tools, but this is not proven by the current frontend flow.',
    ['Finals Guide, page 4: three minutes, five minutes of questions, jury device','apps/web/src/state/finn.ts: match and playFixture','apps/web/src/state/finn-scenarios.ts: d1October and d1November'])

s=slide('Trust needs a reason.','Atlas · October 2026')
rect(s,0.73,2.54,4.32,3.72,NAVY,None,True)
text(s,'22 October',1.04,2.97,3.75,0.9,38,IVORY,True)
text(s,'2026',1.06,3.9,3.7,0.48,25,'D6DEE8')
text(s,'Approved Atlas exception\nBelgium · October only',1.07,4.78,3.65,0.94,21,IVORY)
for yy,code,head,body,col in [(2.62,'S4','Approved client agreement','Roy approved the Atlas exception.',GREEN),(3.88,'S1','General procedure','The client exception takes precedence.',MUTED),(5.14,'S10 / S3','Teams proposals and recollections','The 23rd and 25th are not approvals.',AMBER)]:
    text(s,code,5.65,yy,1.3,0.42,19,col,True)
    text(s,head,7.0,yy,5.6,0.48,22,NAVY,True)
    text(s,body,7.0,yy+0.51,5.6,0.63,19,MUTED)
notes(s,'Why 22 October applies',20,
    'The account owner approved the Atlas agreement for Belgium in October. It overrides the general deadline. The other dates are proposals or recollections. Finn makes that distinction visible, with the source behind the answer.',
    'Use this slide as the first demo fallback. Avoid reading a confidence percentage as calibrated accuracy. The fixture contains 0.92, but this is a scripted score, not an evaluation result.',
    'Fixture reasoning and UI source display are present in inspected code. Trust scores and applicability decisions are authored fixture values, not independently computed or validated predictions.',
    ['apps/web/src/state/finn-scenarios.ts: d1October, loonmutatieSources','apps/web/src/components/FinnResultContent.tsx: answer sources and Why this source'])

s=slide('No evidence? Make the gap visible.','Atlas · November 2026')
text(s,'“And for next month?”',0.75,2.55,11.7,0.7,31,MUTED)
rect(s,0.74,3.54,5.59,2.23,WHITE,BORDER,True)
text(s,'No substantiated answer',1.04,3.88,4.94,0.66,27,NAVY,True)
text(s,'October’s agreement does not\nconfirm November’s deadline.',1.04,4.73,4.94,0.78,22,MUTED)
rect(s,6.65,3.54,5.92,2.23,PALE,None,True)
text(s,'Next step: Roy',6.97,3.88,5.18,0.66,27,NAVY,True)
text(s,'Ask the Atlas account owner\nfor a confirmed November date.',6.97,4.73,5.18,0.78,22,MUTED)
text(s,'Contact is simulated in this prototype. No message is sent.',0.76,6.23,11.72,0.42,18,AMBER,True)
notes(s,'The right next step',15,
    'For November, Finn does not recycle October’s answer. It shows the missing evidence and identifies Roy, the Atlas owner. Contact is simulated here. A real message and reply loop comes next.',
    'Use this slide as the second demo fallback. Do not say that Finn has actually sent a Teams message. In the fixture, all October sources are no longer applicable to November, including the general rule.',
    'The null answer, owner and proposed message are fixture values. The UI explicitly says simulated contact. No external message has been sent.',
    ['apps/web/src/state/finn-scenarios.ts: d1November, nextStep','apps/web/src/components/FinnResultContent.tsx: Simulated contact in this demo'])

s=slide('Same principle. Another HR question.','Sick leave example')
text(s,'When must a sick leave notice be passed on?',0.75,2.37,11.7,0.52,25,MUTED)
for x,big,tag,body,bg,col in [(0.74,'24 hours','S6 · Approved','Current Belgian\nprocedure',LIGHT_GREEN,GREEN),(4.98,'48 hours','S7 · Unconfirmed','A Teams message\nwithout approval',AMBER_BG,AMBER),(9.22,'72 hours','S34 · Expired','Older procedure\nvalid until 2025',PALE,MUTED)]:
    rect(s,x,3.26,3.39,2.52,bg,None,True)
    text(s,big,x+0.24,3.54,2.95,0.65,31,col,True)
    text(s,tag,x+0.24,4.42,2.95,0.37,17,NAVY,True)
    text(s,body,x+0.24,4.98,2.95,0.67,19,MUTED)
text(s,'Synthetic demo policy · this is not payroll or legal advice',0.76,6.32,11.72,0.32,15,MUTED)
notes(s,'Sick leave example',15,
    'In our synthetic sick leave case, the approved procedure says twenty four hours. Forty eight is unconfirmed. Seventy two comes from an expired version. The status of the source matters.',
    'Optional cut if the live demo runs long. These are synthetic fixture policy values, not real Belgian legal requirements. Do not present them as guidance to SD Worx clients.',
    'd2 fixture content only. Not an independent policy, legal or payroll validation. The fixture score 0.85 is authored, not calibrated.',
    ['apps/web/src/state/finn-scenarios.ts: d2, S6, S7, S34, S32'])

s=slide('Less searching. Better judgement.','Value for SD Worx',True)
text(s,'Help people use the right context\nwithout maintaining another knowledge portal.',0.74,2.46,11.6,0.97,28,IVORY)
rect(s,0.76,3.91,11.77,0.025,'49617A')
text(s,'NEXT STEP · A FOCUSED PILOT',0.77,4.29,11.4,0.3,13,'F3BA55',True)
for x,n,head,body in [(0.77,'01','Connect safely','Authorised sources\nand access controls'),(4.94,'02','Close the loop','Real colleague confirmation\nwhen evidence is missing'),(9.08,'03','Measure value','Time to a supported answer\nand expert interruptions')]:
    text(s,n,x,4.9,0.55,0.37,16,CORAL,True)
    text(s,head,x,5.4,3.55,0.43,22,IVORY,True)
    text(s,body,x,5.94,3.65,0.68,17,'D6DEE8')
notes(s,'Value and next step',25,
    'For SD Worx, everyday expertise becomes easier to use, with uncertainty in view. Next, we would pilot authorised sources and real colleague confirmation in one workflow, measuring time to a supported answer and interruptions to experts. Less searching, with better judgement.',
    'End here within 170 seconds. Leave ten seconds of safety margin. These are proposed pilot metrics, not measured savings. The safe connection and access controls are future pilot requirements, not a claim of a completed security review. Keep the completed Aikido evidence ready separately for submission and questions.',
    'Value hypothesis and roadmap. Real connectors, permission enforcement over enterprise sources, ongoing freshness monitoring, colleague replies and knowledge reuse remain future work. Voice is deliberately excluded. No security scan result is asserted.',
    ['Finals Guide, pages 6, 10, 11 and 12: criteria, challenge and Aikido requirements','README.md: challenge coverage, privacy, stretch goals and source choices'])

draft = TMP/'SDtrust_Finals_Draft_v1_base.pptx'
prs.save(draft)
(OUT/'slide_content.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
lines=['# SDtrust Finals Draft v1', '',
    'Eight slides. Planned duration 170 seconds including a 45 second demo, with ten seconds of margin below the guide’s three minute ceiling. Five minutes of jury questions are separate.', '',
    'The main stage agenda lists a ten minute finalist slot but does not establish a separate speaking allowance. This draft follows the explicit three minute pitch rule on pages 4 and 6. Expand only after organiser confirmation.', '',
    'English copy follows British spelling. Text, source panels and diagrams are editable. Finn is approved raster artwork. The logo is the official SVG. The demo frame is deliberately a placeholder, not a screenshot.', '',
    '## Slide outline and first speaking script', '']
for m in meta:
    lines.extend([f"### {m['slide']}. {m['title']} · {m['seconds']} seconds",'',m['spoken'],'','Presenter cues: '+m['cues'],'','Evidence boundary: '+m['status'],'','Sources: '+'; '.join(m['sources']),''])
lines.extend(['## Before the final handoff','',
    '1. Replace slide 4 with a capture of the final app version, or keep slides 5 and 6 as the fallback. Capture the actual supported interaction and keep the synthetic data label.',
    '2. Reconfirm the delivered app behaviour with the application coordinator. This presentation session read code and existing evidence only. It did not operate or build the app.',
    '3. Rehearse on the jury device, including any recording. No recording or live hyperlink is embedded in v1.',
    '4. If source search or contact becomes real, update slide 6 and the spoken boundary only after verification. Do not add voice claims just because voice work exists in another session.',
    '5. Prepare the final Aikido scan screenshots separately. No security result or performance improvement has been measured or claimed in this deck.',
    '6. The generic docs/task-brief.md is an empty template. It supplied no extra timing, acceptance criteria or runtime evidence.',
    '7. Commit and push remain with the coordinating thread. This session has made no Git mutations.', '',
    '## Suggested jury answers','',
    'What works today? The inspected frontend contains the three prepared fixture flows, source explanations, a missing evidence state, and a simulated next step. This session did not rerun the app.', '',
    'Is the trust score an accuracy metric? No. Scores and reasons are authored fixture content. A production trust method would need evaluation and calibration.', '',
    'Does Finn contact a colleague? No external message is sent in this prototype. The UI shows the intended handoff using synthetic data.', '',
    'How will knowledge stay fresh? The prototype demonstrates refusing an answer when a source no longer applies. A pilot would add authorised live sources, freshness checks and confirmation reuse.', '',
    'What is the security position? The deck uses synthetic records and contains no credentials. Enterprise source permissions and final Aikido findings must be assessed separately. The deck does not certify security.', ''])
(OUT/'SDtrust_Finals_Draft_v1_Notes.md').write_text('\n'.join(lines),encoding='utf-8')
source_files=list(sources.values())+['README.md','AGENTS.md','docs/task-brief.md','apps/web/src/desktop.css','apps/web/src/styles.css','apps/web/src/state/finn-scenarios.ts','apps/web/src/state/finn.ts','apps/web/src/components/FinnResultContent.tsx','docs/design/finn_pet_refined/session_status.json','docs/design/desktop_concepts/prompts.json']
manifest=[]
for source in source_files:
    b=(ROOT/source).read_bytes()
    manifest.append({'source':source,'sha256':hashlib.sha256(b).hexdigest(),'bytes':len(b)})
guide=Path('C:/Users/Sebas/Downloads/Tectonic Hackathon - Finals Guide-1.pdf')
manifest.append({'source':str(guide),'sha256':hashlib.sha256(guide.read_bytes()).hexdigest(),'pages_used':[4,6,10,11,12,16]})
(OUT/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'slides':len(prs.slides),'seconds':sum(m['seconds'] for m in meta),'script_words':sum(len(m['spoken'].split()) for m in meta),'base':str(draft)}))
