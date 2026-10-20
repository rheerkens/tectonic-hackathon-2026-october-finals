"""Build the refined v3 pitch base. Writes only to docs/presentation/v3 and .local/presentation/v3.

Content and notes come from the reviewed v1 slide_content.json (read only).
The official SVG logo is placed later by render_v3.ps1 at each 'LOGO_SLOT' shape.
"""
from pathlib import Path
import sys, json, re

ROOT = Path(__file__).resolve().parents[3]
V1 = ROOT / 'docs/presentation'
OUT = V1 / 'v3'
TMP = ROOT / '.local/presentation/v3'
TMP.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(ROOT / '.local/presentation/vendor'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE
from PIL import Image
from lxml import etree as E

A = OUT / 'assets'
# Brand tokens: apps/web/src/desktop.css plus the brief's coral and amber.
NAVY = '192D46'; NAVY2 = '29435F'; IVORY = 'F8F6F1'; SURFACE = 'FFFEFD'; SURFACE2 = 'F1EEE7'
CORAL = 'EE6558'; AMBER = 'F3BA55'; MUTED = '5F6F85'; BORDER = 'E5E1D8'; BORDER2 = 'C9CED6'
PALE = 'E9EEF3'; ON_DARK = 'D6DEE8'; RULE_DARK = '49617A'
OK = '0F7A3A'; OK_BG = 'EAF3EC'; WARN = 'A76400'; WARN_BG = 'FFF0D7'

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
cp = prs.core_properties
cp.title = 'SDtrust | Tectonic Hackathon Finals | Dashboard v3'
cp.subject = 'Dashboard visual version of the reviewed v1 pitch, with an editable demo storyboard'
cp.author = 'Sebastian, Roy and Wanne'
cp.keywords = 'SDtrust, Finn, SD Worx, Tectonic, v3'
cp.comments = 'Synthetic fixture examples. This deck does not certify app runtime or security.'

content = json.loads((V1 / 'slide_content.json').read_text(encoding='utf-8'))
TOTAL = len(content)

def rgb(h): return RGBColor.from_string(h)

NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
def ae(name, **attrs):
    return E.Element('{'+NS+'}'+name, **{k:str(v) for k,v in attrs.items()})

def gradient(q, colors, radial=False, angle=45):
    """Native DrawingML gradient. Colours are (hex, opacity from 0 to 1)."""
    sp=q._element.spPr
    for node in list(sp):
        if E.QName(node).localname in ['solidFill','noFill','gradFill','blipFill','pattFill']:
            sp.remove(node)
    g=ae('gradFill',rotWithShape='1'); stops=ae('gsLst')
    for i,(color,alpha) in enumerate(colors):
        stop=ae('gs',pos=round(i*100000/(len(colors)-1)))
        c=ae('srgbClr',val=color);c.append(ae('alpha',val=round(alpha*100000)))
        stop.append(c);stops.append(stop)
    g.append(stops)
    if radial:
        path=ae('path',path='circle');path.append(ae('fillToRect',l=50000,t=50000,r=50000,b=50000));g.append(path)
    else: g.append(ae('lin',ang=round(angle*60000),scaled='1'))
    geometry=next((n for n in sp if E.QName(n).localname in ['prstGeom','custGeom']),None)
    sp.insert(list(sp).index(geometry)+1 if geometry is not None else 1,g)

def glow(s,x,y,w,h,color,opacity):
    q=box(s,x,y,w,h,None,shape=MSO_SHAPE.OVAL,name='Dashboard ambient '+color)
    gradient(q,[(color,0),(color,opacity)],radial=True)
    for prior in list(q._element.spPr):
        if E.QName(prior).localname in ['effectLst','effectDag']: q._element.spPr.remove(prior)
    effects=ae('effectLst');effects.append(ae('softEdge',rad=508000));q._element.spPr.append(effects)
    return q

def glass(s,x,y,w,h,dark=False,name='Dashboard glass panel',radius=.22):
    q=box(s,x,y,w,h,None,'31475D' if dark else 'FFFFFF',radius=radius,name=name)
    gradient(q,[('29435F',.97),('172637',.95)] if dark else [('FFFFFF',.90),('FFFEFD',.68)],angle=105)
    effect=ae('effectLst'); shadow=ae('outerShdw',blurRad=177800,dist=44450,dir=5400000,algn='ctr',rotWithShape='0')
    c=ae('srgbClr',val='000000' if dark else NAVY);c.append(ae('alpha',val=16000 if dark else 6500));shadow.append(c);effect.append(shadow)
    sp=q._element.spPr
    for prior in list(sp):
        if E.QName(prior).localname in ['effectLst','effectDag']: sp.remove(prior)
    sp.append(effect)
    return q

def aurora_text(shape):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            rpr=run._r.get_or_add_rPr()
            for child in list(rpr):
                if E.QName(child).localname in ['solidFill','gradFill']: rpr.remove(child)
            g=ae('gradFill',rotWithShape='1'); gs=ae('gsLst')
            for pos,color in [(0,'D84935'),(55000,'C26D12'),(100000,'4A74C0')]:
                stop=ae('gs',pos=pos);stop.append(ae('srgbClr',val=color));gs.append(stop)
            g.append(gs);g.append(ae('lin',ang=0,scaled='1'))
            latin=next((n for n in rpr if E.QName(n).localname=='latin'),None)
            rpr.insert(list(rpr).index(latin) if latin is not None else 0,g)

def icon_tile(s,x,y,kind,color='coral'):
    q=box(s,x,y,.44,.44,None,radius=.11,name='Dashboard icon tile '+kind)
    gradient(q,[('D84935',1),('EE6558',1)] if color=='coral' else [('4A74C0',1),('6F9BE0',1)],angle=45)
    # Native line icon, inspired by the dashboard row symbols.
    if kind=='source':
        box(s,x+.135,y+.09,.17,.25,None,'FFFFFF',radius=.02,lw=1.25)
        rule(s,x+.17,y+.17,x+.275,y+.17,'FFFFFF',1)
        rule(s,x+.17,y+.23,x+.275,y+.23,'FFFFFF',1)
    else:
        rule(s,x+.11,y+.22,x+.195,y+.30,'FFFFFF',1.6)
        rule(s,x+.195,y+.30,x+.34,y+.13,'FFFFFF',1.6)

def box(s, x, y, w, h, fill=None, line=None, radius=None, shape=None, dash=False, lw=1.0, name=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    q = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        q.adjustments[0] = radius / min(w, h)
    if fill: q.fill.solid(); q.fill.fore_color.rgb = rgb(fill)
    else: q.fill.background()
    if line:
        q.line.color.rgb = rgb(line); q.line.width = Pt(lw)
        if dash: q.line.dash_style = MSO_LINE.DASH
    else: q.line.fill.background()
    q.shadow.inherit = False
    if name: q.name = name
    return q

def rule(s, x1, y1, x2, y2, color, w=1.0, name='Rule'):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w); c.name = name
    st = c._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: c._element.remove(st)  # drop theme effect (shadow) on rules
    return c

def text(s, paras, x, y, w, h, size=18, color=NAVY, bold=False, align=None, spacing=1.05,
         anchor=MSO_ANCHOR.TOP, spc=0, name=None):
    """paras: str (newline = paragraph) or list of paragraphs, each a list of (text, opts)."""
    if isinstance(paras, str):
        paras = [[(line, {})] for line in paras.split('\n')]
    q = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = q.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    flat = []
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = spacing
        if align is not None: p.alignment = align
        for t, o in runs:
            assert not re.search('[‐-―-]', t), repr(t)  # no hyphens or dashes in visible prose
            r = p.add_run(); r.text = t; flat.append(t)
            f = r.font; f.name = 'Manrope'; f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold); f.color.rgb = rgb(o.get('color', color))
            rpr = r._r.get_or_add_rPr()
            if o.get('strike'): rpr.set('strike', 'sngStrike')
            if o.get('spc', spc): rpr.set('spc', str(int(o.get('spc', spc) * 100)))
    q.name = name or ('Text: ' + ' '.join(flat)[:60])
    return q

def chip(s, label, x, y, fill, color, w=None, size=12, h=0.36, line=None, name=None):
    w = w or (0.3 + len(label) * size * 0.0098)
    box(s, x, y, w, h, fill, line, radius=h / 2, name=name or ('Chip: ' + label))
    text(s, label, x, y, w, h, size, color, True, PP_ALIGN.CENTER, 1.0, MSO_ANCHOR.MIDDLE, name='Chip text: ' + label)
    return w

def picture(s, file, x, y, w, h, name):
    iw, ih = Image.open(file).size
    sc = min(w / iw, h / ih); dw, dh = iw * sc, ih * sc
    q = s.shapes.add_picture(str(file), Inches(x + (w - dw) / 2), Inches(y + h - dh), Inches(dw), Inches(dh))
    q.name = name
    return q

def logo_slot(s, x, y, d):
    q = box(s, x, y, d, d, name='LOGO_SLOT')  # replaced by official SVG in render_v3.ps1
    return q

def new_slide(dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY if dark else IVORY)
    # The dashboard body has three restrained radial colour washes.
    glow(s,8.7,0,4.63,3.8,'EE6558',.18 if dark else .13)
    glow(s,0,3.8,5.4,3.7,'6F9BE0',.17 if dark else .13)
    glow(s,7.3,3.3,4.0,3.7,'F2A55A',.08)
    return s

def chrome(s, kicker, title, dark=False, title_size=40):
    n = len(prs.slides)
    glass(s,.53,.31,12.25,.75,dark,name='Dashboard brand strip',radius=.18)
    logo_slot(s, 0.7, 0.5, 0.34)
    text(s, 'SDtrust', 1.12, 0.52, 1.6, 0.32, 15, IVORY if dark else NAVY, True, name='Wordmark')
    text(s, f'{n:02} / {TOTAL:02}', 11.33, 0.56, 1.3, 0.26, 11, ON_DARK if dark else MUTED, True,
         PP_ALIGN.RIGHT, spc=1, name='Page number')
    text(s, kicker.upper(), 0.7, 1.24, 9.0, 0.26, 12, AMBER if dark else CORAL, True, spc=2, name='Kicker')
    if title:
        text(s, title, 0.7, 1.56, 11.9, 0.7, title_size, IVORY if dark else NAVY, True, spacing=1.0, name='Title')

def notes(s, i):
    m = content[i]
    cues = m['cues'].replace('No video or live link is embedded in v1.', 'No video or live link is embedded in this deck.')
    m['cues'] = cues
    body = (f"{m['title']}\nTiming: {m['seconds']} seconds.\n\nSPEAK\n{m['spoken']}\n\nPRESENTER CUES\n{cues}"
            f"\n\nEVIDENCE BOUNDARY\n{m['status']}\n\nSOURCES\n" + '\n'.join(m['sources']))
    s.notes_slide.notes_text_frame.text = body

# 1 · Cover: editorial split, Finn as host
s = new_slide()
box(s, 8.45, 0, 4.883333, 7.5, NAVY, name='Cover panel')
glow(s,8.45,.2,4.8,4.8,'EE6558',.30)
glow(s,8.45,3.0,4.8,4.5,'6F9BE0',.28)
glass(s,8.78,1.28,4.0,5.72,True,'Finn glass companion panel',radius=.34)
picture(s, A / 'Finn.png', 8.85, 1.25, 4.1, 4.75, 'Approved asset: Finn original')
text(s, 'Finn', 8.95, 6.28, 3.9, 0.34, 16, IVORY, True, PP_ALIGN.CENTER, name='Finn name')
text(s, 'Your SDtrust desktop companion', 8.95, 6.62, 3.9, 0.3, 12, ON_DARK, False, PP_ALIGN.CENTER, name='Finn role')
logo_slot(s, 0.7, 0.62, 0.46)
text(s, 'SDtrust', 1.28, 0.66, 2.2, 0.4, 20, NAVY, True, name='Wordmark')
text(s, 'TECTONIC HACKATHON FINALS · SD WORX', 0.72, 2.0, 7.4, 0.28, 12, CORAL, True, spc=2, name='Kicker')
text(s, [[('The answer.', {})], [('And the right', {})]],
     0.68, 2.38, 7.6, 2.05, 62, NAVY, True, spacing=0.94, name='Cover headline')
accent=text(s,'next step.',.68,4.13,7.6,1.1,62,CORAL,True,spacing=.94,name='Aurora headline')
aurora_text(accent)
text(s, 'Meet Finn, your SDtrust knowledge companion.', 0.72, 5.6, 7.4, 0.45, 21, MUTED, name='Cover subline')
rule(s, 0.72, 6.42, 1.32, 6.42, NAVY, 1.25)
text(s, 'Sebastian · Roy · Wanne', 0.72, 6.58, 6.0, 0.36, 15, NAVY, True, name='Team')
notes(s, 0)

# 2 · Problem: four equally loud numbers
s = new_slide(); chrome(s, 'The problem', 'One client. Four different deadlines.')
text(s, '“When can Atlas submit payroll changes?”', 0.7, 2.42, 11.5, 0.42, 22, MUTED, name='Question')
cols = [(0.7, 0.00, '20', 'S1 · PROCEDURE', 'General rule', 'Belgium'),
        (3.78, 0.42, '22', 'S4 · AGREEMENT', 'Client agreement', 'Atlas'),
        (6.86, 0.14, '23', 'S10 · TEAMS', 'Meeting proposal', 'Teams'),
        (9.94, 0.52, '25', 'S3 · TEAMS', 'Chat recollection', 'Teams')]
for x, off, big, tag, head, sub in cols:
    glass(s,x-.13,2.99,2.94,3.32,name='Dashboard source panel '+big)
    y = 3.1 + off
    chip(s, tag, x, y, SURFACE2, MUTED, size=11, h=0.32)
    text(s, big, x - 0.04, y + 0.32, 2.7, 1.8, 112, NAVY, True, spacing=0.9, name='Numeral ' + big)
    text(s, head, x, y + 1.95, 2.8, 0.36, 17, NAVY, True)
    text(s, sub, x, y + 2.29, 2.8, 0.3, 13, MUTED)
text(s, 'October 2026', 11.33, 2.5, 1.3, 0.26, 12, MUTED, False, PP_ALIGN.RIGHT, name='Month')
text(s, [[('Finding a message is easy. ', {}), ('Knowing what applies is the hard part.', {'color': CORAL})]],
     0.7, 6.42, 11.9, 0.42, 22, NAVY, True, name='Problem statement')
notes(s, 1)

# 3 · Meet Finn: guide plus three step flow
s = new_slide(); chrome(s, 'Meet Finn', 'Ask once. See the answer and why.')
glass(s,.95,2.75,3.85,3.85,name='Companion glass stage',radius=.6)
glow(s,1.0,2.9,3.6,3.4,'F2A55A',.16)
picture(s, A / 'Finn_wave.png', 1.15, 2.85, 3.45, 3.55, 'Approved asset: Finn waving body_lock_03')
steps = [('1', 'Ask in your own words', 'One small companion on your desktop.', NAVY),
         ('2', 'Check the evidence', 'The answer, its sources and the reasons.', NAVY),
         ('3', 'Know what comes next', 'If evidence is missing, surface the owner.', CORAL)]
for k in range(3): glass(s,5.62,2.66+k*1.25,7.08,1.09,name='Dashboard action row')
rule(s, 6.13, 3.15, 6.13, 5.65, BORDER2, 1.5, 'Step connector')
for k, (n, head, body, col) in enumerate(steps):
    y = 2.83 + k * 1.25
    stepdot=box(s, 5.85, y, 0.56, 0.56, col, radius=.16, name='Step dot ' + n)
    gradient(stepdot,[('D84935',1),('EE6558',1)] if k==2 else [('4A74C0',1),('6F9BE0',1)])
    text(s, n, 5.85, y, 0.56, 0.56, 18, IVORY, True, PP_ALIGN.CENTER, 1.0, MSO_ANCHOR.MIDDLE)
    text(s, head, 6.8, y - 0.02, 5.8, 0.4, 25, NAVY, True)
    text(s, body, 6.8, y + 0.44, 5.8, 0.34, 17, MUTED)
notes(s, 2)

# 4 · Demo: stage moment, storyboard placeholder
s = new_slide(True); chrome(s, 'Product demo', 'Let’s ask Finn.', dark=True)
glass(s,.6,2.4,8.45,4.1,True,'Demo outer bezel',radius=.24)
glass(s, 0.7, 2.5, 8.25, 3.9, name='Demo frame',radius=.16)
for k in range(3):
    box(s, 1.0 + k * 0.22, 2.74, 0.12, 0.12, BORDER2, shape=MSO_SHAPE.OVAL, name='Frame dot')
rule(s, 0.7, 3.06, 8.95, 3.06, BORDER, 1.0, 'Frame bar')
chip(s, 'DEMO PLACEHOLDER', 1.05, 3.4, NAVY, IVORY, size=10, h=0.34)
text(s, 'Final app capture goes here', 1.05, 4.05, 7.5, 0.55, 30, NAVY, True)
text(s, '45 seconds · ask, inspect, follow up', 1.05, 4.7, 7.5, 0.34, 18, MUTED)
text(s, 'Editable storyboard. This is not an app screenshot.', 1.05, 5.8, 7.5, 0.26, 12, MUTED)
rule(s, 9.62, 2.72, 9.62, 5.98, RULE_DARK, 1.25, 'Timeline')
beats = [('0:00', 'Ask', 'Atlas deadline?'), ('0:10', 'Answer', '22 October'),
         ('0:25', 'Evidence', 'S4 and one rejected source'), ('0:35', 'Follow up', 'Next month?')]
for k, (t, head, body) in enumerate(beats):
    y = 2.55 + k * 0.98
    box(s, 9.54, y + 0.1, 0.16, 0.16, CORAL, shape=MSO_SHAPE.OVAL, name='Beat dot')
    text(s, t, 9.92, y, 0.75, 0.3, 12, AMBER, True, spc=1)
    text(s, head, 10.62, y - 0.03, 2.0, 0.34, 18, IVORY, True)
    text(s, body, 10.62, y + 0.33, 2.01, 0.55, 13, ON_DARK)
text(s, 'Scripted fixture demo · synthetic records', 0.7, 6.68, 8.2, 0.28, 13, ON_DARK, name='Demo label')
notes(s, 3)

# 5 · The answer vs the noise
s = new_slide(); chrome(s, 'Atlas · October 2026', 'Trust needs a reason.')
answer=glass(s, 0.7, 2.45, 4.75, 4.3, True, name='Answer card',radius=.22)
glow(s,.85,2.6,4.35,3.9,'6F9BE0',.15)
glass(s,5.78,2.45,7.0,4.3,name='Dashboard evidence panel')
icon_tile(s,1.05,2.76,'confirmed','blue')
text(s, 'THE ANSWER', 1.68, 2.86, 3.0, 0.26, 11, AMBER, True, spc=2)
text(s, '22 October', 1.05, 3.12, 4.2, 0.75, 46, IVORY, True, spacing=1.0)
text(s, '2026', 1.07, 3.9, 2.0, 0.36, 22, ON_DARK)
text(s, 'Approved exception\nAtlas · Belgium\nOctober only', 1.07, 4.75, 2.55, 1.0, 16, IVORY, spacing=1.12)
chip(s, 'Source S4', 1.07, 6.0, NAVY2, IVORY, size=11, h=0.34)
picture(s, A / 'Finn_review.png', 3.4, 4.25, 1.95, 2.42, 'Approved asset: Finn review approved02')
rows = [('22', False, OK, 'S4', 'Approved client agreement', 'Roy approved the Atlas exception.', 'Applies', OK_BG, OK),
        ('20', True, MUTED, 'S1', 'General procedure', 'The client exception takes precedence.', 'Overridden', SURFACE2, MUTED),
        ('23 25', True, MUTED, 'S10 / S3', 'Teams proposals', 'A proposed 23rd and a recalled 25th. Neither is an approval.', 'Not approved', WARN_BG, WARN)]
for k, (d, strike, dcol, code, head, body, status, sbg, scol) in enumerate(rows):
    y = 2.5 + k * 1.43
    if k: rule(s, 6.0, y - 0.2, 12.63, y - 0.2, BORDER, 1.0)
    text(s, [[(d, {'strike': strike})]], 6.0, y - 0.1, 1.55, 0.8, 40, dcol, True, spacing=1.0, name='Date ' + d)
    text(s, code, 7.6, y, 1.6, 0.24, 11, MUTED, True, spc=1)
    text(s, head, 7.6, y + 0.26, 3.5, 0.5, 18, NAVY, True)
    text(s, body, 7.6, y + 0.66, 3.5, 0.65, 14, MUTED, spacing=1.1)
    chip(s, status, 11.23, y + 0.02, sbg, scol, w=1.4, size=11, h=0.34)
notes(s, 4)

# 6 · Gap to person
s = new_slide(); chrome(s, 'Atlas · November 2026', 'No evidence? Make the gap visible.')
text(s, '“And for next month?”', 0.7, 2.42, 11.5, 0.42, 22, MUTED, name='Question')
glass(s,0.7,3.1,4.95,3.0,name='Evidence glass panel')
box(s, 0.7, 3.1, 4.95, 3.0, None, BORDER2, radius=0.18, dash=True, lw=1.5, name='Evidence gap')
text(s, '?', 1.05, 3.15, 1.2, 1.45, 80, BORDER2, True, spacing=1.0, name='Gap mark')
text(s, 'No substantiated answer', 1.05, 4.5, 4.4, 0.4, 23, NAVY, True)
text(s, 'October’s agreement does not\nconfirm November’s deadline.', 1.05, 5.0, 4.4, 0.7, 16, MUTED, spacing=1.12)
picture(s, A / 'Finn_waiting.png', 5.68, 3.55, 2.15, 2.55, 'Approved asset: Finn waiting drawn_gesture_03')
glass(s, 7.9, 3.1, 4.73, 3.0, True, name='Next step card',radius=.22)
q=box(s, 8.25, 3.45, 0.78, 0.78, CORAL, radius=.23, name='Roy avatar')
gradient(q,[('D84935',1),('EE6558',1)])
text(s, 'R', 8.25, 3.45, 0.78, 0.78, 26, IVORY, True, PP_ALIGN.CENTER, 1.0, MSO_ANCHOR.MIDDLE, name='Roy initial')
text(s, 'NEXT STEP', 9.25, 3.5, 3.0, 0.26, 11, AMBER, True, spc=2)
text(s, 'Roy', 9.25, 3.76, 3.0, 0.5, 30, IVORY, True, spacing=1.0)
text(s, 'Atlas account owner', 8.25, 4.55, 4.1, 0.3, 15, ON_DARK)
text(s, 'Ask for a confirmed\nNovember date.', 8.25, 5.0, 4.1, 0.7, 18, IVORY, True, spacing=1.12)
box(s, 0.7, 6.42, 11.93, 0.5, WARN_BG, radius=0.25, name='Simulation banner')
text(s, 'Contact is simulated in this prototype. No message is sent.', 1.0, 6.42, 11.4, 0.5, 15, WARN, True,
     anchor=MSO_ANCHOR.MIDDLE, name='Simulation note')
notes(s, 5)

# 7 · Same principle: status drives emphasis
s = new_slide(); chrome(s, 'Sick leave example', 'Same principle. Another HR question.')
text(s, '“When must a sick leave notice be passed on?”', 0.7, 2.42, 11.5, 0.42, 22, MUTED, name='Question')
nums = [(0.7, 4.5, '24', 132, NAVY, False, 'S6 · Approved', OK_BG, OK, 'Current Belgian procedure'),
        (5.5, 3.6, '48', 84, WARN, False, 'S7 · Unconfirmed', WARN_BG, WARN, 'A Teams message without approval'),
        (9.4, 3.23, '72', 60, '6F7E91', True, 'S34 · Expired', SURFACE2, MUTED, 'Older procedure, valid until 2025')]
for x, w, n, size, col, strike, tag, bg, tc, body in nums:
    glass(s,x-.14,3.05,w-.02,3.5,name='Dashboard policy panel '+n)
    text(s, [[(n, {'strike': strike}), (' hours', {'size': max(18, size // 4), 'strike': False})]],
         x - 0.03, 2.95, w, 2.3, size, col, True, spacing=0.9, anchor=MSO_ANCHOR.BOTTOM, name='Hours ' + n)
    chip(s, tag, x, 5.48, bg, tc, size=12, h=0.36)
    text(s, body, x, 6.0, w - 0.2, 0.6, 15, MUTED)
text(s, 'Synthetic demo policy · this is not payroll or legal advice', 0.7, 6.72, 11.9, 0.28, 13, MUTED, name='Synthetic note')
notes(s, 6)

# 8 · Close
s = new_slide(True); chrome(s, 'Value for SD Worx', None, dark=True)
text(s, 'Less searching.', 0.7, 1.58, 11.9, .94, 54, IVORY, True, spacing=.98, name='Title')
text(s,'Better judgement.',.7,2.48,11.9,.94,54,IVORY,True,spacing=.98,name='Closing promise')
text(s, 'Help people use the right context without maintaining another knowledge portal.',
     0.7, 3.5, 8.6, 0.8, 21, ON_DARK, spacing=1.12, name='Value line')
text(s, 'NEXT STEP · A FOCUSED PILOT', 0.7, 4.97, 6.0, 0.26, 12, AMBER, True, spc=2)
chip(s, 'Future pilot, not yet measured', 9.83, 4.92, None, ON_DARK, w=2.8, size=11, h=0.34, line=RULE_DARK)
for k in range(3): glass(s,.6+k*4.2,5.29,4.12,1.55,True,'Pilot glass panel',radius=.19)
pilot = [('1', 'Connect safely', 'Authorised sources\nand access controls'),
         ('2', 'Close the loop', 'Real colleague confirmation\nwhen evidence is missing'),
         ('3', 'Measure value', 'Time to a supported answer\nand expert interruptions')]
for k, (n, head, body) in enumerate(pilot):
    x = 0.7 + k * 4.2
    box(s, x, 5.45, 0.56, 0.56, CORAL if k == 0 else NAVY2, CORAL if k == 0 else RULE_DARK,
        shape=MSO_SHAPE.OVAL, name='Pilot dot ' + n)
    text(s, n, x, 5.45, 0.56, 0.56, 17, IVORY, True, PP_ALIGN.CENTER, 1.0, MSO_ANCHOR.MIDDLE)
    text(s, head, x + 0.75, 5.47, 3.2, 0.4, 20, IVORY, True)
    text(s, body, x + 0.75, 5.92, 3.3, 0.66, 14, ON_DARK, spacing=1.12)
notes(s, 7)

base = TMP / 'SDtrust_Finals_Dashboard_v3_base.pptx'
prs.save(base)

# English notes companion
lines = ['# SDtrust Finals Dashboard v3 · speaker notes', '',
         'Eight slides. Planned duration 170 seconds including a 45 second demo, ten seconds below the explicit three minute pitch ceiling (Finals Guide pages 4 and 6). Five minutes of jury questions are separate. Criteria and challenge framing: Finals Guide pages 6 and 10.', '',
         'This is the visual refinement of the reviewed v1 draft. Script, cues, evidence boundaries and sources are carried over unchanged, except that the slide 4 cue now says no video or live link is embedded in this deck. Each note keeps three kinds of evidence apart: code evidence that was inspected, the scripted fixture demo, and the roadmap.', '',
         'British spelling. All text, diagrams and the demo storyboard are editable PowerPoint objects. The logo is the official SVG; Finn uses the original artwork and approved frames (waving body_lock_03, review approved02, waiting drawn_gesture_03).', '',
         '## Slides', '']
for m in content:
    lines += [f"### {m['slide']}. {m['title']} · {m['seconds']} seconds", '', m['spoken'], '',
              'Presenter cues: ' + m['cues'], '', 'Evidence boundary: ' + m['status'], '',
              'Sources: ' + '; '.join(m['sources']), '']
lines += ['## Before presenting', '',
          '1. Slide 4 is a deliberate, labelled placeholder. Replace it with a capture of the final verified app, or use slides 5 and 6 as the fallback walkthrough. Keep the synthetic data label.',
          '2. Atlas and sick leave figures are synthetic demo fixtures, not Belgian law or SD Worx policy. Trust scores in the fixtures are authored values, not accuracy measures.',
          '3. Contact is simulated; no message is sent. Do not claim real connector search, real colleague contact or new voice features.',
          '4. Slide 8 describes a proposed pilot. Its measures are hypotheses, not achieved savings.',
          '5. Rehearse timing and test on the jury device. Neither was verified when building this deck.', '']
(OUT / 'SDtrust_Finals_Dashboard_v3_Notes.md').write_text('\n'.join(lines), encoding='utf-8')
print(json.dumps({'slides': len(prs.slides), 'seconds': sum(m['seconds'] for m in content), 'base': str(base)}))
